#!/usr/bin/env python3
"""Backup Quimicer - Etapa 7: gera as planilhas FR 001, FR 002 e FR 011 a partir do banco,
no mesmo layout das planilhas originais (igual ao botao "Exportar" do app).
Entrada: pasta com itens.json, registros.json, lotes.json, config.json (exportados via psql).
Uso: python3 gerar_excel.py <pasta_json> <pasta_saida>"""
import json, sys, os, re, math, unicodedata, datetime as dt
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

SRC, OUT = sys.argv[1], sys.argv[2]
HOJE = dt.datetime.now(ZoneInfo('America/Sao_Paulo')).date()
MESES = ['Janeiro','Fevereiro','Março','Abril','Maio','Junho','Julho','Agosto','Setembro','Outubro','Novembro','Dezembro']

F_DENS = dict(k='dens', nome='Densidade', unit='g/cm³', tipo='faixa', mn='densMin', mx='densMax')
F_PH = dict(k='ph', nome='pH', tipo='faixa', mn='phMin', mx='phMax')
F_VISC = dict(k='visc', nome='Viscosidade', unit='s', tipo='faixa', mn='viscMin', mx='viscMax')
MODS = {
  'mp': dict(key='mp', fr='FR 001', rev='00', title='Liberação de Matéria-Prima', itemCol='Matéria-prima', aux=('fornecedor','Fornecedor'), cfg='mp', validade=True, laudo=True,
    fields=[F_DENS, F_PH, F_VISC, dict(k='colagem', nome='Colagem', tipo='aprov', txt='colagem'), dict(k='aparencia', nome='Aparência', unit='visual', tipo='aprov', txt='visual'), dict(k='odor', nome='Odor', tipo='aprov', txt='odor')]),
  'pa': dict(key='pa', fr='FR 002', rev='00', title='Liberação de Produto Acabado', itemCol='Produto', aux=('op','Ordem de produção'), cfg='geral', validade=False, laudo=False,
    fields=[F_DENS, F_PH, F_VISC, dict(k='umid', nome='Umidade', unit='%', tipo='livre'), dict(k='imperm', nome='Impermeabilidade', tipo='aprov', txt='diluicao'),
            dict(k='colagem', nome='Colagem', tipo='aprov', txt='colagem'), dict(k='aparencia', nome='Aparência', unit='visual', tipo='aprov', txt='visual')]),
  'pi': dict(key='pi', fr='FR 011', rev='01', title='Liberação de Produto Intermediário', itemCol='Intermediário', aux=('op','Ordem de produção'), cfg='pi', validade=False, laudo=False,
    fields=[F_DENS, F_PH, F_VISC, dict(k='umid', nome='Umidade', unit='%', tipo='faixa', mn='umidMin', mx='umidMax'), dict(k='residuo', nome='Resíduo', unit='%', tipo='faixa', mn='resMin', mx='resMax'),
            dict(k='colagem', nome='Colagem', tipo='aprov', txt='colagem'), dict(k='defloc', nome='Defloculação', tipo='aprov', txt='defloculacao'), dict(k='aparencia', nome='Aparência', unit='visual', tipo='aprov', txt='visual'),
            dict(k='pesoIni', nome='Peso inicial', unit='kg', tipo='livre'), dict(k='pesoFin', nome='Peso final', unit='kg', tipo='livre'), dict(k='perda', nome='Perda na reação', unit='%', tipo='calc')]),
}
AV = {'ok': 'CONFORME', 'nc': 'N/C', 'na': 'N/A', 'pend': 'NÃO AVALIADO', 'free': ''}
VSTAT = {'venc': 'Vencida', 'prox': 'Vence em breve', 'ok': 'Dentro da validade', 'lib': 'Liberada pela direção', 'sem': 'Sem validade', 'enc': 'Inativo'}

def load(n):
    with open(os.path.join(SRC, n), encoding='utf-8') as f: return json.load(f) or []
def slug(s):
    s = unicodedata.normalize('NFKD', s); s = ''.join(c for c in s if not unicodedata.combining(c))
    return (re.sub(r'^-|-$', '', re.sub(r'[^A-Z0-9]+', '-', s.upper()))[:120]) or 'X'
def normNome(s): return re.sub(r'\s+', ' ', str(s or '').strip()).upper()
def parseNum(v):
    if v is None or isinstance(v, bool): return None
    if isinstance(v, (int, float)): return float(v)
    t = str(v).strip().replace(',', '.')
    try: n = float(t) if t else None
    except ValueError: return None
    return n if n is not None and math.isfinite(n) else None
def fmt(n):
    if n is None or n == '': return '—'
    if isinstance(n, (int, float)) and not isinstance(n, bool):
        s = ('%.4f' % n).rstrip('0').rstrip('.'); return s.replace('.', ',')
    return str(n)
def check(v, mn, mx):
    if mn is None or mx is None: return 'na' if v is None or v == '' else 'free'
    n = parseNum(v)
    if n is None: return 'pend'
    return 'ok' if mn - 1e-9 <= n <= mx + 1e-9 else 'nc'
def checkDens(v, mn, mx):
    b = check(v, mn, mx)
    if b != 'nc': return b
    n = parseNum(v); dec = lambda x: len(repr(float(x)).split('.')[1].rstrip('0')) if float(x) != int(float(x)) else 0
    d = max(dec(mn), dec(mx), 2); m = 10 ** d
    tr = math.floor(n * m + 1e-9) / m; ro = math.floor(n * m + 1e-7 + 0.5) / m
    ok = lambda x: mn - 1e-9 <= x <= mx + 1e-9
    return 'ok' if ok(tr) or ok(ro) else 'nc'
def textCheck(v):
    if v is None or v == '': return 'na'
    t = str(v).upper()
    return 'nc' if re.search(r'N[ÃA]O|NOK|REPROV', t) else ('ok' if t == 'OK' else 'free')
def perda(r):
    a, b = parseNum(r.get('pesoIni')), parseNum(r.get('pesoFin'))
    return round((a - b) / a * 10000) / 100 if a is not None and b is not None and a > 0 else None
def dateOf(s):
    try: return dt.date.fromisoformat(str(s)[:10])
    except Exception: return None

itens, regs, lotes, cfgs = load('itens.json'), load('registros.json'), load('lotes.json'), load('config.json')
CFG = {c['chave']: c['dados'] or {} for c in cfgs}
FRREV = CFG.get('frrev', {})
ITENS = {}
for it in itens: ITENS.setdefault(it['modulo'], {})[it['id']] = it['dados'] or {}
LOTES = {l['id']: l['dados'] or {} for l in lotes}

def specOf(rec, mod):
    p = ITENS.get(mod['key'], {}).get(slug(normNome(rec.get('produto'))))
    if p:
        o = {'extras': p.get('extras') or []}
        for f in mod['fields']:
            if f['tipo'] == 'faixa': o[f['mn']] = p.get(f['mn']); o[f['mx']] = p.get(f['mx'])
        return o, p
    return rec.get('spec') or {}, {}

def evalRec(rec, mod):
    sp, item = specOf(rec, mod); r = {'x': {}}; fora = []
    for f in mod['fields']:
        v = rec.get(f['k'])
        if f['tipo'] == 'faixa': s = (checkDens if f['k'] == 'dens' else check)(v, sp.get(f['mn']), sp.get(f['mx']))
        elif f['tipo'] in ('livre', 'calc'): s = 'na' if v is None or v == '' else 'free'
        else: s = (('pend' if f.get('txt') and item.get(f['txt']) else 'na') if v is None or v == '' else textCheck(v))
        r[f['k']] = s
        if s == 'nc': fora.append(f['nome'])
    for e in (sp.get('extras') or []) if isinstance(sp.get('extras'), list) else []:
        v = (rec.get('extras') or {}).get(e.get('id'))
        s = (textCheck(v) if v else 'pend') if e.get('tipo') == 'aprovacao' else check(v, e.get('min'), e.get('max'))
        r['x'][e.get('id')] = s
        if s == 'nc': fora.append(e.get('nome') or e.get('id'))
    if mod['validade'] and rec.get('validade') and rec.get('data') and rec['validade'] < rec['data']: fora.append('Validade')
    r['fora'] = fora; return r, sp

def lotInfo(rec):
    L = LOTES.get(re.sub(r'[^A-Za-z0-9_\-.~:@+]', '_', rec['_id']), {})
    val = L.get('validade') or rec.get('validade') or ''; d = L.get('direcao')
    vd = dateOf(val)
    if L.get('encerrado'): st = 'enc'
    elif not vd: st = 'sem'
    elif vd < HOJE: st = 'lib' if d and (not d.get('ate') or (dateOf(d['ate']) or HOJE) >= HOJE) else 'venc'
    else: st = 'prox' if (vd - HOJE).days <= 30 else 'ok'
    return vd, VSTAT[st]

def recs_by_month(mod):
    meses = {}
    rows = sorted([g for g in regs if g['modulo'] == mod['key']], key=lambda g: (g['mes'], g['origem'] != 'planilha', g['ordem'] if g['ordem'] is not None else 0))
    idx = {}
    for g in rows:
        r = dict(g['dados'] or {}); m = g['mes']
        if g['origem'] == 'planilha':
            i = idx.get(m, 0); idx[m] = i + 1; r['_src'] = 'arquivo'; r['_id'] = f'{m}:a{i}'
        else: r['_src'] = 'live'; r['_id'] = f"{m}:{g['id']}"
        if mod['key'] == 'pi' and r.get('perda') is None: r['perda'] = perda(r)
        meses.setdefault(m, []).append(r)
    return meses

# ---------- layout igual às planilhas originais FR 001 / FR 002 / FR 011 (mesmo do botão Exportar do app) ----------
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
AZ, AM, VD = 'FFDBEEF4', 'FFFFFF00', 'FFD4EA6B'
THIN = Side(style='thin', color='FF000000'); BORDA = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
def solid(c): return PatternFill('solid', fgColor=c)
def txt(v): return '-' if v is None or v == '' else v
def nz(v): return 'N/A' if v is None or v == '' else v
def faixa(a, b): return '-' if a is None or b is None else f'{fmt(a)} a {fmt(b)}'
def tx(v): return '-' if v is None or str(v).strip() == '' else str(v)
def numlote(v):
    if v is None or v == '': return ''
    s = str(v).strip()
    return int(s) if re.fullmatch(r'\d{1,15}', s) and not s.startswith('0') else s
def st(v, mn, mx):
    """mesmo resultado da fórmula IF(AND(v>=min,v<=max),"CONFORME",IF(min="N/A","N/A","N/C")) — no Excel texto > número"""
    val, a, b = txt(v), nz(mn), nz(mx)
    def ge(x, y): return (x >= y if isinstance(x, (int, float)) else str(x).upper() >= str(y).upper()) if isinstance(x, (int, float)) == isinstance(y, (int, float)) else isinstance(x, str)
    def le(x, y): return (x <= y if isinstance(x, (int, float)) else str(x).upper() <= str(y).upper()) if isinstance(x, (int, float)) == isinstance(y, (int, float)) else isinstance(x, (int, float))
    return 'CONFORME' if ge(val, a) and le(val, b) else ('N/A' if a == 'N/A' else 'N/C')
def qtdTxt(v, un):
    n = parseNum(str(v).replace('.', '').replace(',', '.')) if isinstance(v, str) and re.fullmatch(r'\d{1,3}(\.\d{3})+(,\d+)?', v.strip()) else parseNum(v)
    return f"{v}" if n is None else f"{n:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.') + ' ' + (un or 'kg')

def layout(mod):
    k = mod['key']
    obs = lambda r, extra='': ' · '.join(x for x in [r.get('obs') or '', extra] if x)
    sit = lambda r, c: r.get('situacao') or ''; rnc = lambda r, c: r.get('rnc') or ''; rncs = lambda r, c: r.get('rncStatus') or ''; ana = lambda r, c: r.get('analista') or ''
    if k == 'mp':
        def obs_mp(r, c):
            vd, _ = lotInfo(r)
            return obs(r, ' · '.join(x for x in ['Fornecedor: ' + r['fornecedor'] if r.get('fornecedor') else '', 'Validade: ' + vd.strftime('%d/%m/%Y') if vd else ''] if x))
        return dict(titulo=f"Liberação de Matéria Prima Fornecedor - {mod['fr']} - Rev{mod['_rev']}", tl=2, hdr=4, halt=33.95, zoom=80, escala=80, altlin=17.85,
            alturas={1: 9.95, 2: 24.95, 3: 9.95}, aba=lambda m: f"{MESES[int(m[5:7]) - 1].upper()}_{m[:4]}", hist='MATERIA_PRIMA', cols=[
            ('Data', 14.43, dict(v=lambda r, c: r['_dt'], fmt='mm-dd-yy')), ('Matéria Prima', 27.14, dict(v=lambda r, c: r.get('produto'))), ('Lote', 17, dict(v=lambda r, c: '' if r.get('lote') is None else str(r.get('lote')), fmt='@')),
            ('Densidade (g/cm³)', 14, dict(wrap=True, v=lambda r, c: txt(r.get('dens')))), ('Ph', 9.57, dict(v=lambda r, c: txt(r.get('ph')))), ('Colagem', 11.57, dict(v=lambda r, c: txt(r.get('colagem')))),
            ('Viscosidade', 13.14, dict(v=lambda r, c: txt(r.get('visc')))), ('Aparência (Visual)', 13, dict(wrap=True, sz=11, v=lambda r, c: txt(r.get('aparencia')))), ('Odor', 13, dict(v=lambda r, c: txt(r.get('odor')))),
            ('Situação', 13, dict(sz=11, v=sit)), ('Nº RNC', 11.29, dict(hfill='FFFF0000', hcor='FFFFFFFF', v=rnc)), ('Status da RNC', 13, dict(hfill='FF0070C0', hcor='FFFFFFFF', wrap=True, v=rncs)),
            ('Analisado por:', 13.14, dict(wrap=True, sz=11, v=ana)), ('Observação', 36.57, dict(v=obs_mp)),
            ('Laudos PDF', 21.14, dict(v=lambda r, c: (r.get('laudo') or {}).get('nome') if isinstance(r.get('laudo'), dict) else (r.get('laudoRef') or '')))])
    if k == 'pi':
        return dict(titulo=f"Liberação de Produto Intermediário - {mod['fr']} – Rev{mod['_rev']}", tl=2, hdr=4, halt=50.1, zoom=80, escala=70, unico='D',
            alturas={1: 9.95, 2: 24.95, 3: 9.95}, aba=lambda m: f"{MESES[int(m[5:7]) - 1].upper()} {m[:4]}", hist='HISTÓRICO',
            cf=[('R', 'cell', [('greaterThanOrEqual', '10', 'FFFF0000'), ('equal', '0', 'FFCC0000')])], cols=[
            ('Data', 14.43, dict(v=lambda r, c: r['_dt'], fmt='mm-dd-yy')), ('Matéria Prima', 35.14, dict(v=lambda r, c: r.get('produto'))), ('Lote', 19.86, dict(v=lambda r, c: numlote(r.get('lote')))),
            ('Ordem de Produção (Nº)', 14, dict(wrap=True, sz=11, v=lambda r, c: txt(r.get('op')))), ('Densidade (g/cm³)', 16.71, dict(wrap=True, v=lambda r, c: txt(r.get('dens')))),
            ('', 8.29, dict(wrap=True, hide=True, v=lambda r, c: '')),
            ('Dm', 13, dict(wrap=True, sz=11, fill=AM, f=lambda n, r: f'VLOOKUP(B{n},bancoquimicer,2,0)', v=lambda r, c: txt(nz(c.get('densMin'))))),
            ('DM', 13, dict(wrap=True, sz=11, fill=AM, f=lambda n, r: f'VLOOKUP(B{n},bancoquimicer,3,0)', v=lambda r, c: txt(nz(c.get('densMax'))))),
            ('status', 11.29, dict(wrap=True, sz=11, fill=AM, hide=True, f=lambda n, r: f'IF(AND(E{n}>=G{n},E{n}<=H{n}),"CONFORME",IF(G{n}="N/A","N/A","N/C"))', v=lambda r, c: st(r.get('dens'), c.get('densMin'), c.get('densMax')))),
            ('Ph', 8.29, dict(v=lambda r, c: txt(r.get('ph')))), ('', 14.57, dict(hide=True, v=lambda r, c: '')), ('', 13, dict(v=lambda r, c: '')),
            ('Phm', 13, dict(wrap=True, sz=11, fill=AM, f=lambda n, r: f'VLOOKUP(B{n},bancoquimicer,5,0)', v=lambda r, c: txt(nz(c.get('phMin'))))),
            ('PHM', 13, dict(wrap=True, sz=11, fill=AM, f=lambda n, r: f'VLOOKUP(B{n},bancoquimicer,6,0)', v=lambda r, c: txt(nz(c.get('phMax'))))),
            ('status', 13, dict(wrap=True, sz=11, fill=AM, f=lambda n, r: f'IF(AND(J{n}>=M{n},J{n}<=N{n}),"CONFORME",IF(M{n}="N/A","N/A","N/C"))', v=lambda r, c: st(r.get('ph'), c.get('phMin'), c.get('phMax')))),
            ('Aparência (Visual)', 14.57, dict(wrap=True, sz=11, v=lambda r, c: txt(r.get('aparencia')))), ('Teste Colagem', 13.29, dict(wrap=True, v=lambda r, c: txt(r.get('colagem')))),
            ('Umidade', 13.14, dict(v=lambda r, c: txt(r.get('umid')))), ('Defloculação', 14.43, dict(v=lambda r, c: txt(r.get('defloc')))),
            ('Resíduo (%) *Copiar valores realizados no processo', 15.86, dict(hfill='FF31859C', hcor='FFFFFFFF', wrap=True, v=lambda r, c: '' if r.get('residuo') in (None, '') else r.get('residuo'))),
            ('Situação', 13.14, dict(sz=11, v=sit)), ('Nº RNC', 8.57, dict(hfill='FFFF0000', hcor='FFFFFFFF', v=rnc)), ('Status da RNC', 6.29, dict(hfill='FF0070C0', hcor='FFFFFFFF', wrap=True, v=rncs)),
            ('Analisado por:', 13.14, dict(wrap=True, sz=11, v=ana)), ('Observação', 31.86, dict(v=lambda r, c: obs(r, 'Viscosidade: ' + fmt(r.get('visc')) if r.get('visc') not in (None, '') else ''))),
            ('Peso Final', 13.86, dict(v=lambda r, c: '' if r.get('pesoFin') in (None, '') else r.get('pesoFin'))), ('Peso Inicial', 11.43, dict(v=lambda r, c: '' if r.get('pesoIni') in (None, '') else r.get('pesoIni'))),
            ('Perda', 15.57, dict(fmt='0.00', f=lambda n, r: f'100-(Z{n}/AA{n}*100)' if perda(r) is not None else None, v=lambda r, c: '' if perda(r) is None else perda(r)))])
    def obs_pa(r, c):
        x = ['Cliente: ' + r['cliente'] if r.get('cliente') else ('Estoque' if r.get('destino') == 'estoque' else ''), 'Qtd: ' + qtdTxt(r['qtdCliente'], r.get('qtdUn')) if r.get('qtdCliente') else '']
        return obs(r, ' · '.join(i for i in x if i))
    return dict(titulo=f"Liberação de Produto Acabado - {mod['fr']} - Rev{mod['_rev']}", tl=3, hdr=6, halt=31.5, zoom=90, escala=80, unico='D', validadata=True, sz=11,
        alturas={1: 9.95, 2: 8.1, 3: 16.35, 4: 16.35}, aba=lambda m: f"{MESES[int(m[5:7]) - 1].upper()} {m[:4]}", hist='HISTÓRICO',
        cf=[('I', 'expr', 'L{r}="N/C"', 'FFFF0000')], cols=[
        ('Data', 14.43, dict(hfill='FFD7E4BD', v=lambda r, c: r['_dt'], fmt='mm-dd-yy')), ('Produto Acabado', 45.86, dict(v=lambda r, c: r.get('produto'))), ('Lote', 19.86, dict(v=lambda r, c: numlote(r.get('lote')))),
        ('Ordem de Produção (Nº)', 13, dict(wrap=True, v=lambda r, c: txt(r.get('op')))),
        ('', 6.57, dict(hfill=AM, fill=AM, wrap=True, hide=True, f=lambda n, r: f'VLOOKUP(B{n},Tabeladados,3,0)', v=lambda r, c: txt(nz(c.get('densMin'))))),
        ('', 13, dict(hfill=AM, fill=AM, wrap=True, f=lambda n, r: f'VLOOKUP(B{n},Tabeladados,4,0)', v=lambda r, c: txt(nz(c.get('densMax'))))),
        ('', 21.14, dict(hfill=AM, fill=AM, wrap=True, hide=True, f=lambda n, r: f'IF(AND(H{n}>=E{n},H{n}<=F{n}),"CONFORME",IF(E{n}="N/A","N/A","N/C"))', v=lambda r, c: st(r.get('dens'), c.get('densMin'), c.get('densMax')))),
        ('Densidade (g/cm³)', 21.14, dict(wrap=True, fill='FFFFFFFF', fmt='0.00', v=lambda r, c: txt(r.get('dens')))), ('Ph', 8.29, dict(v=lambda r, c: txt(r.get('ph')))),
        ('ph minimo', 16.57, dict(hfill=AM, fill=AM, hide=True, f=lambda n, r: f'VLOOKUP($B{n},Tabeladados,7,0)', v=lambda r, c: txt(nz(c.get('phMin'))))),
        ('ph máximo', 13, dict(hfill=AM, fill=AM, f=lambda n, r: f'VLOOKUP($B{n},Tabeladados,8,0)', v=lambda r, c: txt(nz(c.get('phMax'))))),
        ('', 13, dict(hfill=AM, fill=AM, f=lambda n, r: f'IF(AND(I{n}>=J{n},I{n}<=K{n}),"CONFORME",IF(J{n}="N/A","N/A","N/C"))', v=lambda r, c: st(r.get('ph'), c.get('phMin'), c.get('phMax')))),
        ('Viscosidade (s)', 16.57, dict(v=lambda r, c: txt(r.get('visc')))), ('Teste  Impermeabilidade', 14.86, dict(v=lambda r, c: txt(r.get('imperm')))), ('Umidade', 13.14, dict(v=lambda r, c: txt(r.get('umid')))),
        ('Teste de colagem', 13, dict(v=lambda r, c: txt(r.get('colagem')))),
        ('', 13.14, dict(hfill=VD, fill=VD, hide=True, f=lambda n, r: f'VLOOKUP($B{n},Tabeladados,11,0)', v=lambda r, c: txt(nz(c.get('viscMin'))))),
        ('', 13, dict(hfill=VD, fill=VD, f=lambda n, r: f'VLOOKUP($B{n},Tabeladados,12,0)', v=lambda r, c: txt(nz(c.get('viscMax'))))),
        ('', 13, dict(hfill=VD, fill=VD, f=lambda n, r: f'IF(AND(M{n}>=Q{n},M{n}<=R{n}),"CONFORME",IF(Q{n}="N/A","N/A","N/C"))', v=lambda r, c: st(r.get('visc'), c.get('viscMin'), c.get('viscMax')))),
        ('Aparência (Visual)', 16, dict(wrap=True, v=lambda r, c: txt(r.get('aparencia')))), ('Situação', 13.14, dict(v=sit)),
        ('Nº RNC', 11.29, dict(hfill='FFFF0000', hcor='FFFFFFFF', v=rnc)), ('Status da RNC', 13, dict(hfill='FF0070C0', hcor='FFFFFFFF', wrap=True, v=rncs)),
        ('Analisado por:', 13.14, dict(wrap=True, v=ana)), ('Observação', 50, dict(v=obs_pa))])

def dados(wb, mod):
    k = mod['key']; its = sorted([p for p in ITENS.get(k, {}).values() if p.get('ativo') is not False], key=lambda p: normNome(p.get('nome')))
    cab = 'FF215968'
    if k == 'pa':
        nome = 'DADOS_01'; faixa_nome = 'Tabeladados'
        cols = [('PRODUTO', 47.3), ('DENSIDADE  (g/mL)', 13.3), ('DENSIDADE  (g/mL) minima', 18.1), ('DENSIDADE  (g/mL) máxima', 13), ('VISCOSIDADE (s)', 23.7), ('p.H', 16.6), ('p.H minimo', 19.6), ('p.H máximo', 16), ('ANÁLISE VISUAL', 41.6), ('TESTE DE APLICAÇÃO (COLAGEM)', 22.4), ('VISCOSIDADE (s)MINIMA', 22.9), ('VISCOSIDADE (s)MAXIMA', 13), ('umidade', 22.3)]
        linhas = [[p.get('nome'), faixa(p.get('densMin'), p.get('densMax')), nz(p.get('densMin')), nz(p.get('densMax')), faixa(p.get('viscMin'), p.get('viscMax')) + (' ' + p['viscCopo'] if p.get('viscCopo') else ''), faixa(p.get('phMin'), p.get('phMax')), nz(p.get('phMin')), nz(p.get('phMax')), tx(p.get('visual')), tx(p.get('colagem')), nz(p.get('viscMin')), nz(p.get('viscMax')), tx(p.get('umidade'))] for p in its]
    elif k == 'pi':
        nome = 'DADOS'; faixa_nome = 'bancoquimicer'
        cols = [('Código MP', 15.6), ('Produto Intermediário         (Nome Usual)', 34.4), ('DENS. MIN', 13.4), ('DENS. MAX', 13), ('Densidade (g/mL)', 13.4), ('PH MIN', 13.4), ('PH MAX', 13), ('p.H', 13.4), ('Viscosidade', 13.6), ('Aparência (visual)', 33.9), ('Teste de Aplicação  (Colagem)', 18.7), ('Umidade', 13.4), ('Defloculação', 13), ('Resíduo (%)', 25.7)]
        linhas = [[p.get('codigo') or '', p.get('nome'), nz(p.get('densMin')), nz(p.get('densMax')), faixa(p.get('densMin'), p.get('densMax')), nz(p.get('phMin')), nz(p.get('phMax')), faixa(p.get('phMin'), p.get('phMax')), faixa(p.get('viscMin'), p.get('viscMax')), tx(p.get('visual')), tx(p.get('colagem')), faixa(p.get('umidMin'), p.get('umidMax')) if p.get('umidMin') is not None else tx(p.get('umidade')), tx(p.get('defloculacao')), faixa(p.get('resMin'), p.get('resMax')) if p.get('resMin') is not None else tx(p.get('residuo'))] for p in its]
    else:
        nome = 'DADOS'; faixa_nome = None; cab = 'FFFFFFFF'
        cols = [('Código MP', 11.6), ('Matérias-Primas\n(Nome Usual)', 13.7), ('Colunas1', 23.1), ('Densidade  (g/mL)', 18), ('pH', 10), ('Odor', 13.6), ('\nTeste de Aplicação Colagem', 12.1), ('Aparência (visual)', 46.3), ('Análise  de espuma em proveta', 38.7), ('\nDeflocução', 17.9), ('LAUDO', 13.4), ('TESTE IMPERMEABILIDADE', 29.1), ('VISCOSIDADE', 31.6)]
        linhas = [[numlote(p.get('codigo')) if p.get('codigo') else '', p.get('nome'), '', faixa(p.get('densMin'), p.get('densMax')), faixa(p.get('phMin'), p.get('phMax')), tx(p.get('odor')), tx(p.get('colagem')), tx(p.get('visual')), '-', '-', '-', '-', faixa(p.get('viscMin'), p.get('viscMax'))] for p in its]
    ws = wb.create_sheet(nome, 0); ws.sheet_view.zoomScale = 80; ws.freeze_panes = 'A2'
    for i, (t, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(1, i, t); c.font = Font(name='Calibri', size=11, bold=True, color='FF000000' if cab == 'FFFFFFFF' else 'FFFFFFFF'); c.fill = solid(cab); c.border = BORDA; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 48.75 if k == 'mp' else 36
    for j, l in enumerate(linhas, 2):
        ws.row_dimensions[j].height = 18
        for i, v in enumerate(l, 1):
            c = ws.cell(j, i, None if v == '' else v); c.font = Font(name='Calibri', size=11); c.border = BORDA
            c.alignment = Alignment(horizontal='left' if i == 1 and k == 'pa' else 'center', vertical='center', wrap_text=True)
            if k == 'pi' and i in (3, 4, 6, 7): c.fill = solid(AM)
    ult = max(len(linhas) + 1, 2); L2 = get_column_letter(len(cols))
    ws.auto_filter.ref = f'A1:{L2}{ult}'
    if faixa_nome == 'Tabeladados': wb.defined_names['Tabeladados'] = DefinedName('Tabeladados', attr_text=f"'{nome}'!$A$1:${L2}${ult}")
    if faixa_nome == 'bancoquimicer':
        wb.defined_names['bancoquimicer'] = DefinedName('bancoquimicer', attr_text=f"'{nome}'!$B$1:${L2}${ult}")
        wb.defined_names['bancodados'] = DefinedName('bancodados', attr_text=f"'{nome}'!$B$2:${L2}${ult}")

def folha(wb, nome, L, mod, recs):
    ws = wb.create_sheet(nome[:31]); ws.sheet_view.zoomScale = L['zoom']
    for i, (_, w, o) in enumerate(L['cols'], 1):
        cd = ws.column_dimensions[get_column_letter(i)]; cd.width = w
        if o.get('hide'): cd.hidden = True
    for r, h in L['alturas'].items(): ws.row_dimensions[r].height = h
    t = ws.cell(L['tl'], 1, L['titulo']); t.font = Font(name='Calibri', size=16, bold=True); t.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[L['hdr']].height = L['halt']
    for i, (h, _, o) in enumerate(L['cols'], 1):
        c = ws.cell(L['hdr'], i, h or None); c.font = Font(name='Calibri', size=12, bold=True, color=o.get('hcor', 'FF000000')); c.fill = solid(o.get('hfill', AZ)); c.border = BORDA
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=bool(o.get('wrap')))
    for k, r in enumerate(recs):
        n = L['hdr'] + 1 + k; _, sp = evalRec(r, mod)
        if L.get('altlin'): ws.row_dimensions[n].height = L['altlin']
        for i, (_, _, o) in enumerate(L['cols'], 1):
            v = o['v'](r, sp); fx = o['f'](n, r) if o.get('f') else None
            c = ws.cell(n, i, ('=' + fx) if fx else (None if v in ('', None) else v))
            if o.get('fmt'): c.number_format = o['fmt']
            c.font = Font(name='Calibri', size=o.get('sz', L.get('sz', 12))); c.border = BORDA; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=bool(o.get('wrap')))
            if o.get('fill'): c.fill = solid(o['fill'])
    last = L['hdr'] + max(len(recs), 1)
    for cf in L.get('cf', []):
        ref = f"{cf[0]}{L['hdr'] + 1}:{cf[0]}{last}"
        if cf[1] == 'cell':
            for op, val, cor in cf[2]: ws.conditional_formatting.add(ref, CellIsRule(operator=op, formula=[val], font=Font(color=cor)))
        else: ws.conditional_formatting.add(ref, FormulaRule(formula=[cf[2].replace('{r}', str(L['hdr'] + 1))], font=Font(color=cf[3])))
    fim = last + 200
    if L.get('unico'):
        u = L['unico']; dv = DataValidation(type='custom', formula1=f"COUNTIF({u}:{u},{u}{L['hdr'] + 1})<=1", allow_blank=True, showErrorMessage=True, errorTitle='Ordem de produção repetida', error='Esta ordem de produção já foi lançada nesta aba.')
        dv.add(f"{u}{L['hdr'] + 1}:{u}{fim}"); ws.add_data_validation(dv)
    if L.get('validadata'):
        dv = DataValidation(type='date', operator='greaterThan', formula1='44835', allow_blank=True, showErrorMessage=True, error='Informe uma data válida.'); dv.add(f"A{L['hdr'] + 1}:A{fim}"); ws.add_data_validation(dv)
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = 9; ws.page_setup.scale = L['escala']
    ws.page_margins.left = ws.page_margins.right = 0.118; ws.page_margins.top = ws.page_margins.bottom = 0.709
    ws.print_title_rows = f"{L['hdr']}:{L['hdr']}"

def planilha(mod):
    def chave(r):
        lote = str(r.get('lote') or ''); m = re.match(r'(\D*)(\d*)(.*)', lote)
        return (r.get('data') or '', m.group(1), int(m.group(2)) if m.group(2) else -1, m.group(3))
    mod['_rev'] = (FRREV.get(mod['key']) or {}).get('rev') or mod['rev']
    L = layout(mod); meses = recs_by_month(mod)
    def prep(r): r = dict(r); r['_dt'] = dateOf(r.get('data')); return r
    wb = Workbook(); wb.remove(wb.active)
    dados(wb, mod)
    ano = str(HOJE.year); ms = sorted(meses)
    antigos = [m for m in ms if m[:4] < ano]; atuais = [m for m in ms if m[:4] >= ano]
    total = 0
    if antigos:
        l = sorted([r for m in antigos for r in meses[m]], key=chave); total += len(l); folha(wb, L['hist'], L, mod, [prep(r) for r in l])
    for m in atuais:
        l = sorted(meses[m], key=chave)
        if l: total += len(l); folha(wb, L['aba'](m), L, mod, [prep(r) for r in l])
    if len(wb.worksheets) == 1: folha(wb, L['aba'](HOJE.strftime('%Y-%m')), L, mod, [])
    wb.active = len(wb.worksheets) - 1
    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = L['titulo']; wb.properties.creator = 'Quimicer - backup automático'
    agora = dt.datetime.now(ZoneInfo('America/Sao_Paulo'))
    titulo = {'mp': 'Liberação de Matéria Prima Fornecedor', 'pi': 'Liberação de Produto Intermediário', 'pa': 'Liberação de Produto Acabado'}[mod['key']]
    nome = f"{mod['fr'].replace(' ', '_')}_{titulo} - Rev {mod['_rev']} - {agora.strftime('%d-%m-%Y_%Hh%M')}.xlsx"
    wb.save(os.path.join(OUT, nome)); return nome, total

os.makedirs(OUT, exist_ok=True)
for k in ('mp', 'pa', 'pi'):
    nome, n = planilha(MODS[k]); print(f'{nome}: {n} lotes')
